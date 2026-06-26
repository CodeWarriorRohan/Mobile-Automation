from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "Test Case Summary"

# ── Styles ──────────────────────────────────────────────────────────────────
header_fill  = PatternFill("solid", fgColor="1F4E79")
alt_fill     = PatternFill("solid", fgColor="DEEAF1")
white_fill   = PatternFill("solid", fgColor="FFFFFF")
header_font  = Font(bold=True, color="FFFFFF", size=11)
normal_font  = Font(size=10)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
thin         = Side(style="thin", color="B0C4DE")
border       = Border(left=thin, right=thin, top=thin, bottom=thin)

def style(cell, fill, font, align):
    cell.fill      = fill
    cell.font      = font
    cell.alignment = align
    cell.border    = border

# ── Column widths ────────────────────────────────────────────────────────────
ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 10
ws.column_dimensions["D"].width = 50
ws.column_dimensions["E"].width = 90

# ── Header row ───────────────────────────────────────────────────────────────
ws.append(["#", "Screen / Module", "TC ID", "Test Method", "Functionality Covered"])
for cell in ws[1]:
    style(cell, header_fill, header_font, center_align)
ws.row_dimensions[1].height = 28

# ── Data ─────────────────────────────────────────────────────────────────────
rows = [
    (1,  "Login Screen\ntest_login.py",                          "TC-01",  "test_tc01_empty_credentials",                              "Submits login with both fields empty. Asserts app remains on login page — no navigation."),
    (2,  "Login Screen\ntest_login.py",                          "TC-02",  "test_tc02_valid_username_invalid_password",                 "Valid username + wrong password. Asserts login page still displayed."),
    (3,  "Login Screen\ntest_login.py",                          "TC-03",  "test_tc03_invalid_username_valid_password",                 "Wrong username + valid password. Asserts user not navigated away."),
    (4,  "Login Screen\ntest_login.py",                          "TC-04",  "test_tc04_invalid_credentials",                            "Both invalid username and password. Asserts login page remains displayed."),
    (5,  "Login Screen\ntest_login.py",                          "TC-05",  "test_tc05_valid_credentials",                              "Correct username and password. Asserts app navigates away from login screen — session created."),
    (6,  "My Work Screen\ntest_my_work.py",                      "TC-06",  "test_tc06_header_always_visible",                          "Asserts Cilio logo (top-left) and Profile icon (top-right) are visible at all times."),
    (7,  "My Work Screen\ntest_my_work.py",                      "TC-07",  "test_tc07_user_info_card_always_visible",                  "Asserts user name, role, View All Work button, and My Badge button are visible."),
    (8,  "My Work Screen\ntest_my_work.py",                      "TC-08",  "test_tc08_refresh_button_always_visible",                  "Asserts the Refresh button is visible regardless of job data state."),
    (9,  "My Work Screen\ntest_my_work.py",                      "TC-09",  "test_tc09_all_three_work_cards_visible",                   "Asserts Today, Tomorrow, and Yesterday work cards are all rendered on screen."),
    (10, "My Work Screen\ntest_my_work.py",                      "TC-10",  "test_tc10_bottom_navigation_always_visible",               "Asserts Home, Search, and Schedule bottom nav tabs are always present."),
    (11, "My Work Screen\ntest_my_work.py",                      "TC-11",  "test_tc11_no_job_sections_when_all_zero",                  "[Scenario 1] Asserts no job detail sections appear when all card counts = 0. Skipped if any card has jobs."),
    (12, "My Work Screen\ntest_my_work.py",                      "TC-12",  "test_tc12_no_route_button_when_all_zero",                  "[Scenario 1] Asserts Route button does not appear when all counts = 0."),
    (13, "My Work Screen\ntest_my_work.py",                      "TC-13",  "test_tc13_card_with_jobs_is_visible",                      "[Scenario 2] For each card with count > 0, asserts card is visible and count is a positive integer. Taps first card with jobs."),
    (14, "My Work Screen\ntest_my_work.py",                      "TC-14",  "test_tc14_section_header_renders_only_for_cards_with_jobs","[Scenario 2] Taps all three cards; asserts section header appears only for cards with count > 0."),
    (15, "My Work Screen\ntest_my_work.py",                      "TC-15",  "test_tc15_log_work_card_summary",                         "Reads all card names and counts via card-reader utility. Logs total and per-card breakdown. Asserts each card visible."),
    (16, "My Work Screen\ntest_my_work.py",                      "TC-16",  "test_tc16_pickup_report_visible_when_jobs_exist",          "[Scenario 2] Taps card with jobs, scrolls down. Asserts Pickup Report button is present in DOM."),
    (17, "My Work Screen\ntest_my_work.py",                      "TC-17",  "test_tc17_job_card_name_and_pay_visible",                  "[Scenario 2] Taps card with jobs. Asserts job customer name and crew pay amount are visible."),
    (18, "My Work Screen\ntest_my_work.py",                      "TC-18",  "test_tc18_expand_job_details_shows_all_fields",            "[Scenario 2] Expands each job detail card; asserts all label fields visible; logs field values."),
    (19, "Profile Screen\ntest_profile.py",                      "TC-25",  "test_tc25_profile_screen_elements",                       "Asserts the Profile screen is displayed with all expected elements present."),
    (20, "Profile Screen\ntest_profile.py",                      "TC-26",  "test_tc26_my_badge_navigation",                           "Taps My Badge button, waits for screen to load, taps back arrow to return to Profile."),
    (21, "Profile Screen\ntest_profile.py",                      "TC-27",  "test_tc27_my_account_navigation",                         "Taps My Account, waits, scrolls down and up to verify scrollability, navigates back."),
    (22, "Profile Screen\ntest_profile.py",                      "TC-28",  "test_tc28_view_cilio_dashboard_navigation",               "Taps View Cilio Dashboard to open dashboard or external link; returns via Profile icon."),
    (23, "Profile Screen\ntest_profile.py",                      "TC-29",  "test_tc29_privacy_policy_page",                           "Taps Privacy Policy; page loads; scrolls down and up; navigates back via back arrow."),
    (24, "Profile Screen\ntest_profile.py",                      "TC-30",  "test_tc30_terms_conditions_page",                         "Taps Terms & Conditions; page loads; scrolls down and up; navigates back."),
    (25, "Profile Screen\ntest_profile.py",                      "TC-31",  "test_tc31_logout_functionality",                          "Taps Logout; asserts confirmation dialog appears; taps Cancel — user remains logged in."),
    (26, "Profile Screen\ntest_profile.py",                      "TC-32",  "test_tc32_logout_functionality",                          "Taps Logout; asserts dialog appears; taps Confirm — initiates the logout flow."),
    (27, "Profile Screen\ntest_profile.py",                      "TC-33",  "test_tc33_valid_credentials",                             "Re-authenticates after logout with valid credentials; asserts login succeeds and session is restored."),
    (28, "View All Work Screen\ntest_view_all_work.py",          "TC-40",  "test_tc40_header_always_visible",                         "Asserts Cilio logo and Profile icon visible on View All Work header."),
    (29, "View All Work Screen\ntest_view_all_work.py",          "TC-41",  "test_tc41_user_info_card_always_visible",                 "Asserts user name, role, My Work button, and My Badge button are visible."),
    (30, "View All Work Screen\ntest_view_all_work.py",          "TC-42",  "test_tc42_find_jobs_by_type_cards_visible",               "Reads all job type cards; logs names and counts; scrolls each into view; asserts visibility. Asserts at least one card found."),
    (31, "View All Work Screen\ntest_view_all_work.py",          "TC-43",  "test_tc43_bottom_navigation_always_visible",              "Asserts Home and Search tabs visible; Schedule tab is DISABLED on this screen."),
    (32, "View All Work Screen\ntest_view_all_work.py",          "TC-44",  "test_tc44_tap_profile_icon_navigates_to_profile",         "Taps Profile icon; waits for Profile screen; taps back arrow to return to View All Work."),
    (33, "View All Work Screen\ntest_view_all_work.py",          "TC-45",  "test_tc45_tap_search_nav_navigates_to_search",            "Taps Search bottom nav icon; waits for Search screen to load."),
    (34, "Job Status Screen\ntest_job_status.py",                "TC-45",  "test_tc45_tap_job_type_card_shows_job_status_screen",     "Reads job type cards; taps first with count > 0; asserts Job Status screen header matches card name. Skipped if all counts = 0."),
    (35, "Job Status Screen\ntest_job_status.py",                "TC-46",  "test_tc46_tap_job_status_card_with_count",               "Reads Job Status sub-cards; taps first with count > 0; waits for results screen to load."),
    (36, "Job Status Screen\ntest_job_status.py",                "TC-47",  "test_tc47_tap_job_status_card_with_zero_count",           "Taps a random result card from the currently displayed Job Status results list."),
    (37, "Global Search Screen\ntest_global_search.py",          "TC-45",  "test_tc45_search_screen_loaded",                         "Asserts Cilio logo and search bar are visible on the Search screen."),
    (38, "Global Search Screen\ntest_global_search.py",          "TC-46",  "test_tc46_search_bar_placeholder_visible",               "Asserts 'Search...' placeholder text is visible in the search bar before any input."),
    (39, "Global Search Screen\ntest_global_search.py",          "TC-47",  "test_tc47_find_jobs_by_filters_label_visible",           "Asserts 'Find Jobs by Filters' section label is visible."),
    (40, "Global Search Screen\ntest_global_search.py",          "TC-50",  "test_tc50_bottom_navigation_visible",                    "Asserts Home and Search tabs visible; Schedule tab disabled on Search screen."),
    (41, "Global Search Screen\ntest_global_search.py",          "TC-51",  "test_tc51_tap_search_bar_opens_input",                   "Taps the search bar; waits for the input field to open."),
    (42, "Global Search Screen\ntest_global_search.py",          "TC-52",  "test_tc52_enter_search_text",                            "Types 'testing' into the search bar. Asserts input field text attribute equals 'testing'."),
    (43, "Global Search Screen\ntest_global_search.py",          "TC-55",  "test_tc55_tap_random_three_search_result_cards",         "Taps 3 random result cards from 'testing' search; navigates back after each; logs each card tapped."),
    (44, "Global Search Screen\ntest_global_search.py",          "TC-56",  "test_tc56_tap_clear_search",                             "Taps clear (x) button; waits for search input to clear."),
    (45, "Global Search Screen\ntest_global_search.py",          "TC-57",  "test_tc57_tap_search_back_arrow",                        "Taps back arrow in search input mode; waits for navigation back to Search screen."),
    (46, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-51",  "test_tc51_tap_basic_search_navigates_to_basic_search",   "Taps Basic filter button; waits for Basic Search screen to load."),
    (47, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-52",  "test_tc52_basic_search_fields_and_go",                   "Asserts Go and Clear buttons visible; uses field-reader to log all enabled filter field names for the current user."),
    (48, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-53",  "test_tc53_enter_customer_first_name_and_go",             "Enters 'testing' in Customer First Name field; taps Go; waits for results."),
    (49, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-54",  "test_tc54_tap_random_basic_search_result_cards",         "Taps up to 2 random result cards from Customer First Name search; navigates back after."),
    (50, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-55",  "test_tc55_tap_clear_shows_confirmation_dialog_ok_clear", "Taps Clear; asserts confirmation dialog shown; taps Ok to confirm the clear action."),
    (51, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-56",  "test_tc56_enter_city_and_go",                            "Enters 'city' in City field; taps Go; waits for search results."),
    (52, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-57",  "test_tc57_tap_random_cards_after_city_search",           "Taps up to 2 random result cards from City search; navigates back after."),
    (53, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-58",  "test_tc58_tap_clear_shows_confirmation_dialog_ok_clear", "Taps Clear and confirms dialog to reset City search filters."),
    (54, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-59",  "test_tc59_enter_phone_number_and_go",                    "Enters '1234567890' in Phone Number field; taps Go."),
    (55, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-60",  "test_tc60_tap_random_cards_after_phone_search",          "Taps up to 2 random result cards from Phone Number search; navigates back."),
    (56, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-61",  "test_tc61_tap_clear_shows_confirmation_dialog_ok_clear", "Taps Clear and confirms dialog to reset Phone Number search filters."),
    (57, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-62",  "test_tc62_navigate_back_from_basic_search",              "Taps back arrow; asserts main Search screen is displayed again."),
    (58, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-63",  "test_tc63_tap_advanced_search_navigates_to_advanced_search","Taps Advanced filter button; waits for Advanced Search screen to load."),
    (59, "Basic & Advanced Search\ntest_basic_advance_search.py","TC-64",  "test_tc64_advanced_search_status_and_go",                "Logs all available Job Status options; selects 'New' from the status dropdown; taps Go."),
    (60, "Hybrid WebView\ntest_hybrid_webview.py",               "—",      "test_switch_to_webview",                                 "Asserts NATIVE_APP context exists; switches to WEBVIEW context; switches back; asserts NATIVE_APP restored."),
    (61, "Hybrid WebView\ntest_hybrid_webview.py",               "—",      "test_webview_current_context_is_native_by_default",      "Asserts default context after app launch is NATIVE_APP."),
    (62, "Job Details — Base Page\ntest_base_jobpage.py",        "TC-59",  "test_tc59_job_details_header_visible",                   "Asserts header, Hot Buttons, Measures button, dynamic job name, and Lead Safe button are all visible."),
    (63, "Job Details — Base Page\ntest_base_jobpage.py",        "TC-60",  "test_tc60_hot_buttons_and_navigation",                   "Taps Hot Buttons element; opens Hot Buttons screen; taps back arrow to return to Job Details."),
    (64, "Job Details — Base Page\ntest_base_jobpage.py",        "TC-61",  "test_tc61_measures_button_disabled",                     "Asserts Measures button is visible but its enabled/clickable attribute is false."),
    (65, "Job Details — Base Page\ntest_base_jobpage.py",        "TC-62",  "test_tc62_cust_site_and_job_sched_buttons",              "Asserts Cust/Site, Job/Sched, and Actions/Docs tab buttons are all visible."),
    (66, "Job Details — Base Page\ntest_base_jobpage.py",        "TC-63",  "test_tc63_measures_button_and_navigation",               "Taps Measures button; navigates back via back arrow."),
    (67, "Job Details — Base Page\ntest_base_jobpage.py",        "TC-64",  "test_tc64_tap_profile_icon",                             "Taps Profile icon to navigate to Profile screen; taps back arrow to return to Job Details."),
    (68, "Seller Info Tab\ntest_seller_info.py",                 "TC-65",  "test_tc65_seller_info_tab_visible",                      "Asserts the Seller Info tab is visible in the Cust/Site tab bar of Job Details."),
    (69, "Seller Info Tab\ntest_seller_info.py",                 "TC-66",  "test_tc66_tap_seller_info_tab",                          "Taps Seller Info tab; asserts store name and store address fields are both non-empty."),
    (70, "Site Info Tab\ntest_site_info.py",                     "TC-64",  "test_tc64_site_info_fields_visible",                     "Clicks Site Info tab; asserts Distance to Seller, Avg Time to Seller, Built Pre-1978, and Lead Safe Job fields are visible."),
    (71, "Site Info Tab\ntest_site_info.py",                     "TC-65",  "test_tc65_zillow_button_info",                           "Taps Zillow icon; asserts bottom sheet title, listing data, webview link, and close button are all visible."),
    (72, "Site Info Tab\ntest_site_info.py",                     "TC-66",  "test_tc66_zillow_button_navigation",                     "Taps Zillow webview link inside the bottom sheet; navigates back to Site Info tab via back arrow."),
    (73, "Site Info Tab\ntest_site_info.py",                     "TC-65",  "test_tc65_is_lead_safe_job_checkbox_checked",            "[Conditional — checkbox checked] Asserts Built Pre-1978 dropdown, Lead Safe Practices Required dropdown, Lead Collection Form button, and Responsible User visible. Skipped if checkbox is unchecked."),
    (74, "Site Info Tab\ntest_site_info.py",                     "TC-66",  "test_tc66_is_lead_safe_job_checkbox_unchecked",          "[Conditional — checkbox unchecked] Asserts conditional fields (Lead Safe question, dropdowns) are hidden. Lead Collection Form button and Responsible User still present. Skipped if checkbox is checked."),
    (75, "Site Info Tab\ntest_site_info.py",                     "TC-69",  "test_tc69_tap_lead_collection_form_button",              "Scrolls to bring Lead Collection Form button into view; taps it; navigates back via back arrow."),
    (76, "Site Info Tab\ntest_site_info.py",                     "TC-70",  "test_tc70_tap_responsible_user_icon",                    "Scrolls to Responsible User section; taps icon; asserts bottom sheet title visible; closes the sheet."),
    (77, "Customer Info Tab\ntest_cust_info.py",                 "TC-67",  "test_tc67_scroll_customer_info_page",                    "Scrolls down to 'Customer Contact Info' label then back to top — verifies the full form is scrollable."),
    (78, "Customer Info Tab\ntest_cust_info.py",                 "TC-68",  "test_tc68_edit_customer_info",                           "Taps Edit Info icon; asserts Save Changes and Cancel buttons visible; taps Cancel; re-enters edit mode."),
    (79, "Customer Info Tab\ntest_cust_info.py",                 "TC-69",  "test_tc69_update_and_save_customer_info",                "Verifies Country dropdown parent, clickable descendants, and hint text. Selects a different country. Fills all editable fields (First Name, Last Name, Phone, Alt Phone, Address) and saves."),
    (80, "Related Orders Tab\ntest_related_orders.py",           "TC-70",  "test_tc70_click_related_order_tab",                      "Taps Related Order tab; asserts both Related and Umbrella sub-tabs are visible."),
    (81, "Related Orders Tab\ntest_related_orders.py",           "TC-71",  "test_tc71_click_related_tab",                            "Taps Related sub-tab; if job cards are present, expands first card and logs. Skipped if no related cards found."),
    (82, "Related Orders Tab\ntest_related_orders.py",           "TC-72",  "test_tc72_click_umbrella_tab",                           "Taps Umbrella sub-tab; if job cards are present, expands first card and logs. Skipped if no umbrella cards found."),
    (83, "Job Info Tab\ntest_job_info.py",                       "TC-73",  "test_tc73_click_job_info_tab_verify_fields_visible",     "Clicks Job/Sched tab to set context; then clicks Job Info sub-tab."),
    (84, "Job Info Tab\ntest_job_info.py",                       "TC-74",  "test_tc74_click_job_info_1_tab",                         "Clicks Job Information 1 sub-tab. Scrolls each dropdown into view and selects a random option for Job Status, Labor Category, and Job Type. Asserts each selected value is reflected."),
    (85, "Job Info Tab\ntest_job_info.py",                       "TC-75",  "test_tc75_edit_job_info_1_fields",                       "Clears and fills Job Number, Project Number, and Purchase Order Number with test values. Asserts each field's text matches the entered value."),
    (86, "Attachments Tab\ntest_attachments.py",                 "TC-79",  "test_tc79_click_attachments_docs_tab",                   "Clicks Actions/Docs tab then Attachments sub-tab; verifies the attachment list is loaded."),
    (87, "Attachments Tab\ntest_attachments.py",                 "TC-80",  "test_tc80_scroll_and_click_random_images",               "Scrolls to last attachment; scrolls back to top; taps 3 randomly chosen image attachments; returns to list after each via back arrow."),
    (88, "Attachments Tab\ntest_attachments.py",                 "TC-81",  "test_tc81_tap_three_dots_menu",                          "Taps three-dots menu of a random attachment. Asserts Rename + Remove visible for images; also asserts Edit visible for non-image (PDF/template) files. Dismisses menu."),
    (89, "Attachments Tab\ntest_attachments.py",                 "TC-82",  "test_tc82_rename_attachment",                            "Opens three-dots menu for a random attachment; taps Rename; clears field; types 'Automation rename'; taps Submit."),
    (90, "Attachments Tab\ntest_attachments.py",                 "TC-83",  "test_tc83_edit_attachment",                              "Opens three-dots menu for a non-image (PDF) attachment; taps Edit; waits for PDF viewer to fully load; taps close icon."),
    (91, "Notes & Email Tab\ntest_notes_and_email.py",           "TC-88",  "test_tc88_click_notes_and_email_tab",                    "Clicks Actions/Docs tab then Notes & Email sub-tab to load the notes list."),
    (92, "Notes & Email Tab\ntest_notes_and_email.py",           "TC-89",  "test_tc89_scroll_notes_and_email_list",                  "Scrolls to the end of the Notes & Email list; scrolls back to beginning — verifies full list scrollability."),
    (93, "Notes & Email Tab\ntest_notes_and_email.py",           "TC-90",  "test_tc90_click_add_notes_and_email_button",             "Taps Add Notes button; types 'test note.' in the Note field; taps Add Note to submit and save."),
    (94, "Notes & Email Tab\ntest_notes_and_email.py",           "TC-91",  "test_tc91_tap_filter_icon",                              "Taps filter icon; selects 'Show Standard Note' filter; taps Apply Filters; asserts filter badge count displays as 1; scrolls list."),
    (95, "Notes & Email Tab\ntest_notes_and_email.py",           "TC-92",  "test_tc92_clear_filters",                                "Taps filter icon; taps Clear All; asserts filter badge is no longer visible (filters removed); scrolls through unfiltered list."),
    (96, "Q&A Tab\ntest_q_and_a.py",                            "TC-93",  "test_tc93_tap_q_and_a_tab",                              "Taps Q&A's sub-tab inside Actions/Docs section; waits for Q&A list to load."),
    (97, "Q&A Tab\ntest_q_and_a.py",                            "TC-94",  "test_tc94_open_q_and_a_dropdown",                        "Opens first Q&A category dropdown and logs its name. Selects first available answer option and logs selection. Navigates back."),
    (98, "Activities Tab\ntest_activities.py",                   "TC-93",  "test_tc93_click_activities_tab",                         "Clicks the Activities sub-tab inside the Actions/Docs section to load the Activities list."),
    (99, "Scheduler Screen\ntest_scheduler.py",                  "TC-94",  "test_tc94_schedule_job",                                 "Taps the Schedule tab in bottom navigation to navigate to the Scheduler screen."),
    (100,"Scheduler Screen\ntest_scheduler.py",                  "TC-95",  "test_tc95_header_always_visible",                        "Asserts Cilio logo, Map icon, Profile icon, and Pickup Report button are all visible in the Scheduler header."),
    (101,"Scheduler Screen\ntest_scheduler.py",                  "TC-96",  "test_tc96_bottom_navigation_always_visible",             "Asserts Home, Search, and Schedule bottom nav tabs are all visible on the Scheduler screen."),
    (102,"Scheduler Screen\ntest_scheduler.py",                  "TC-97",  "test_tc_97_tap_calendar_section_items_on_scheduler",     "Taps Left Arrow (previous week), Right Arrow (next week), and Extend/Collapse knob on the calendar control."),
    (103,"Scheduler Screen\ntest_scheduler.py",                  "TC-98",  "test_tc98_tap_map_icon",                                 "Taps Map icon (with coordinate fallback if XPath fails); asserts Route text label and Calendar icon visible on route screen; taps Calendar icon."),
    (104,"Scheduler Screen\ntest_scheduler.py",                  "TC-99",  "test_tc_99_tap_pickup_report",                           "Taps Pickup Report button; interacts with calendar (Left Arrow, Right Arrow, Extend); selects 2-3 consecutive dates from week calendar; navigates back via Android system back gesture."),
    (105,"Scheduler Screen\ntest_scheduler.py",                  "TC-100", "test_tc_100_click_on_the_scheduled_job",                 "Scrolls the scheduler timeline to end then back to beginning; swipes up repeatedly until a job card with AM or PM in content-desc is found; taps a randomly selected job card."),
    (106,"Scheduler Screen\ntest_scheduler.py",                  "TC-101", "test_tc_101_log_bottom_sheet_and_view_job",              "Waits for job bottom sheet to appear; reads and logs all visible text fields (name, address, type, status); asserts at least one detail captured; taps View Job button."),
    (107,"Scheduler Screen\ntest_scheduler.py",                  "TC-102", "test_tc_102_tap_profile_icon_navigates_to_profile",      "Taps Profile icon on Scheduler screen; waits for Profile screen; asserts Profile icon is visible on Profile screen."),
]

for idx, (row_num, screen, tc_id, method, desc) in enumerate(rows, 2):
    ws.append([row_num, screen, tc_id, method, desc])
    row_fill = alt_fill if idx % 2 == 0 else white_fill
    for col in range(1, 6):
        cell = ws.cell(row=idx, column=col)
        style(cell, row_fill, normal_font, center_align if col == 1 else left_align)
    ws.row_dimensions[idx].height = 45

ws.freeze_panes = "A2"

output = "test_case_summary.xlsx"
wb.save(output)
print(f"Done — {output} created with {len(rows)} test cases.")
